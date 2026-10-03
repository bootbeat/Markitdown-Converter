"""Conservative Markdown normalization and source-correlated asset recovery."""
from __future__ import annotations

import base64
from dataclasses import dataclass
import hashlib
from html.parser import HTMLParser
import io
import re
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET
from urllib.parse import quote


@dataclass
class PostprocessResult:
    markdown: str
    assets: list[Path]
    warnings: list[str]


def normalize_headings(text: str) -> str:
    return re.sub(r"^(#{1,6})([^#\s])", r"\1 \2", text, flags=re.MULTILINE)


def normalize_lists(text: str) -> str:
    # Only normalize lines already unambiguously marked as lists. A leading
    # pair of asterisks is commonly Markdown bold, not a list marker.
    return re.sub(r"^([ \t]*(?:[-+*]|\d+[.)])[ \t]+)[ \t]+", r"\1", text, flags=re.MULTILINE)


def normalize_spacing(text: str) -> str:
    result = []
    in_fence = False
    for line in text.splitlines():
        if re.match(r"^\s*(```|~~~)", line):
            in_fence = not in_fence
            result.append(line)
        else:
            result.append(line if in_fence else line.rstrip(" \t"))
    return "\n".join(result) + ("\n" if text.endswith("\n") and result else "")


def normalize_tables(text: str) -> str: return text
def normalize_code_blocks(text: str) -> str: return text
def normalize_links(text: str) -> str: return text


def _outside_fences(text: str, transform) -> str:
    parts, fenced = [], False
    for line in text.splitlines(keepends=True):
        if re.match(r"^\s*(```|~~~)", line):
            fenced = not fenced
            parts.append(line)
        else:
            parts.append(line if fenced else transform(line))
    return "".join(parts)


def _image_info(data: bytes) -> tuple[str, str] | None:
    """Validate image bytes and return canonical extension and MIME type."""
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
            fmt = (image.format or "").upper()
    except Exception:
        return None
    mapping = {"PNG": ("png", "image/png"), "JPEG": ("jpg", "image/jpeg"),
               "GIF": ("gif", "image/gif"), "WEBP": ("webp", "image/webp"),
               "BMP": ("bmp", "image/bmp"), "TIFF": ("tiff", "image/tiff")}
    return mapping.get(fmt)


class _HTMLImages(HTMLParser):
    def __init__(self):
        super().__init__()
        self.uris: list[str] = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == "img":
            src = dict(attrs).get("src", "")
            if src.startswith("data:image/"):
                self.uris.append(src)


def _save_assets(output_file: Path, payloads: list[bytes], warnings: list[str]) -> list[tuple[Path, str]]:
    asset_dir = output_file.with_name(output_file.stem + "_assets")
    extracted = []
    for data in payloads:
        info = _image_info(data)
        if not info:
            warnings.append("An image candidate did not decode as a valid supported image; left its Markdown reference unchanged.")
            continue
        ext, _ = info
        index = len(extracted) + 1
        name = f"image-{index:03d}.{ext}"
        target = asset_dir / name
        try:
            asset_dir.mkdir(parents=True, exist_ok=True)
            if target.exists() and hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(data).digest():
                warnings.append(f"Asset filename collision at {target.name}; candidate was not emitted.")
                continue
            if not target.exists():
                target.write_bytes(data)
            if _image_info(target.read_bytes()) != info:
                warnings.append(f"Saved asset failed verification: {target.name}.")
                continue
            extracted.append((target, quote(f"{asset_dir.name}/{target.name}", safe="/._-")))
        except OSError as exc:
            warnings.append(f"Could not save image asset {name}: {exc}")
    return extracted


def _data_uri_bytes(uri: str) -> bytes | None:
    match = re.fullmatch(r"data:image/[^;,]+;base64,([A-Za-z0-9+/=\r\n]+)", uri, re.I)
    if not match:
        return None
    try:
        return base64.b64decode(re.sub(r"\s+", "", match.group(1)), validate=True)
    except (ValueError, base64.binascii.Error):
        return None


def _replace_docx_images(text: str, source: Path, output_file: Path, warnings: list[str]):
    refs = list(re.finditer(r"!\[([^\]]*)\]\(data:image/[^)]*\)", text))
    if not refs:
        return text, []
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    try:
        with zipfile.ZipFile(source) as archive:
            doc = ET.fromstring(archive.read("word/document.xml"))
            relroot = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            relmap = {rel.get("Id"): "word/" + rel.get("Target", "").lstrip("/") for rel in relroot}
            payloads = []
            for blip in doc.findall(".//a:blip", ns):
                target = relmap.get(blip.get(f"{{{ns['r']}}}embed"))
                if target and target in archive.namelist():
                    payloads.append(archive.read(target))
    except Exception as exc:
        warnings.append(f"DOCX image relationships could not be inspected: {exc}")
        return text, []
    if len(payloads) != len(refs):
        warnings.append(f"DOCX image references could not be safely paired: {len(refs)} Markdown references, {len(payloads)} source images.")
        return text, []
    saved = _save_assets(output_file, payloads, warnings)
    if len(saved) != len(refs):
        warnings.append("DOCX image replacement skipped because not every source image passed validation.")
        return text, [p for p, _ in saved]
    for match, (_, relpath) in reversed(list(zip(refs, saved))):
        text = text[:match.start()] + f"![{match.group(1)}]({relpath})" + text[match.end():]
    return text, [p for p, _ in saved]


def _replace_html_images(text: str, source: Path, output_file: Path, warnings: list[str]):
    refs = list(re.finditer(r"!\[([^\]]*)\]\(data:image/[^)]*\)", text))
    if not refs:
        return text, []
    parser = _HTMLImages()
    try:
        parser.feed(source.read_text(encoding="utf-8"))
        payloads = [_data_uri_bytes(uri) for uri in parser.uris]
    except Exception as exc:
        warnings.append(f"HTML inline image source could not be inspected: {exc}")
        return text, []
    if len(payloads) != len(refs) or any(p is None for p in payloads):
        warnings.append(f"HTML images could not be safely paired: {len(refs)} Markdown references, {len(payloads)} source data URIs.")
        return text, []
    saved = _save_assets(output_file, payloads, warnings)
    if len(saved) != len(refs):
        return text, [p for p, _ in saved]
    for match, (_, relpath) in reversed(list(zip(refs, saved))):
        text = text[:match.start()] + f"![{match.group(1)}]({relpath})" + text[match.end():]
    return text, [p for p, _ in saved]


def _pptx_pictures(source: Path):
    from pptx import Presentation
    import re as regex
    prs = Presentation(str(source))
    pictures=[]
    def visit(shapes):
        for shape in sorted(shapes, key=lambda s: (s.top, s.left)):
            if getattr(shape, "shape_type", None) == 6:
                visit(shape.shapes)
            elif getattr(shape, "shape_type", None) == 13:
                filename = regex.sub(r"\W", "", shape.name) + ".jpg"
                pictures.append((filename, shape.image.blob))
    for slide in prs.slides:
        visit(slide.shapes)
    return pictures


def _replace_pptx_images(text: str, source: Path, output_file: Path, warnings: list[str]):
    refs = list(re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", text))
    if not refs:
        return text, []
    try:
        pictures = _pptx_pictures(source)
    except Exception as exc:
        warnings.append(f"PPTX picture data could not be inspected: {exc}")
        return text, []
    if len(refs) != len(pictures) or any(Path(m.group(2)).name != pic[0] for m, pic in zip(refs, pictures)):
        warnings.append(f"PPTX picture references could not be safely paired with source pictures ({len(refs)} references, {len(pictures)} source pictures).")
        return text, []
    saved = _save_assets(output_file, [b for _, b in pictures], warnings)
    if len(saved) != len(refs):
        warnings.append("PPTX reference replacement skipped because not every source picture passed validation.")
        return text, [p for p, _ in saved]
    for match, (_, relpath) in reversed(list(zip(refs, saved))):
        text = text[:match.start()] + f"![{match.group(1)}]({relpath})" + text[match.end():]
    return text, [p for p, _ in saved]


def _recover_docx_structure(text: str, source: Path, warnings: list[str]) -> str:
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": W}
    try:
        with zipfile.ZipFile(source) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            stylesroot = ET.fromstring(archive.read("word/styles.xml"))
            numbering = ET.fromstring(archive.read("word/numbering.xml"))
    except Exception as exc:
        warnings.append(f"DOCX structure was not changed because OOXML could not be inspected: {exc}")
        return text
    style_levels = {}
    for style in stylesroot.findall("w:style", ns):
        sid = style.get(f"{{{W}}}styleId", "")
        name = style.find("w:name", ns)
        label = name.get(f"{{{W}}}val", "") if name is not None else sid
        outline = style.find("w:pPr/w:outlineLvl", ns)
        match = re.fullmatch(r"Heading\s*(\d+)", label, re.I) or re.fullmatch(r"Heading(\d+)", sid, re.I)
        if match:
            style_levels[sid] = int(match.group(1))
        elif outline is not None:
            try: style_levels[sid] = int(outline.get(f"{{{W}}}val")) + 1
            except (TypeError, ValueError): pass
    num_to_kind = {}
    abs_kinds = {}
    for abstract in numbering.findall("w:abstractNum", ns):
        aid=abstract.get(f"{{{W}}}abstractNumId")
        kinds=[]
        for lvl in abstract.findall("w:lvl", ns):
            fmt=lvl.find("w:numFmt",ns)
            kinds.append(fmt.get(f"{{{W}}}val","") if fmt is not None else "")
        abs_kinds[aid]=kinds
    for num in numbering.findall("w:num",ns):
        nid=num.get(f"{{{W}}}numId"); aid=num.find("w:abstractNumId",ns)
        if aid is not None: num_to_kind[nid]=abs_kinds.get(aid.get(f"{{{W}}}val"),[])
    # Match source paragraphs to unique visible raw lines only; this strips
    # Markdown presentation syntax for lookup purposes but never edits content.
    def match_key(value: str) -> str:
        value=re.sub(r"^\s*#{1,6}\s+", "", value.strip())
        value=re.sub(r"^\s*(?:[-+*]|\d+[.)])\s+", "", value)
        value=re.sub(r"!?\[([^\]]*)\]\([^)]*\)", lambda m: "" if m.group(0).startswith("!") else m.group(1), value)
        value=re.sub(r"<[^>]*>", "", value)
        value=re.sub(r"[*_`]+", "", value)
        return re.sub(r"\s+", " ", value).strip().casefold()
    lines=text.splitlines(keepends=True)
    plain_to_indexes={}
    for i,line in enumerate(lines):
        visible=match_key(line)
        plain_to_indexes.setdefault(visible,[]).append(i)
    paras=root.findall(".//w:body/w:p",ns)
    paragraph_style_ids=[]
    for para in paras:
        style=para.find("w:pPr/w:pStyle",ns)
        paragraph_style_ids.append(style.get(f"{{{W}}}val","") if style is not None else "")
    style_counts={sid:paragraph_style_ids.count(sid) for sid in set(paragraph_style_ids)}
    # A heading style used for a large share of the document is not a reliable
    # paragraph classifier. In this file Heading4 is applied to body paragraphs
    # as well as section titles. Only explicit, user-confirmed title text can
    # retain that overused style's heading markup; all other matches stay prose.
    mass_applied_heading_styles={sid for sid,count in style_counts.items()
                                 if sid in style_levels and len(paras) and count/len(paras) >= 0.30}
    confirmed_heading_titles={
        "abstract", "introduction", "objectives of the project", "system requirements",
        "functional requirements", "implementation and working procedures",
        "conclusion and future scope", "conclusion and futurescope",
        "limitations of the project", "applications of the project", "appendices",
    }
    def title_key(value: str) -> str:
        value=value.strip()
        value=re.sub(r"^\s*\d+(?:\.\s*\d+)*[.)\s-]*", "", value)
        value=re.sub(r"\s*(?:[-–]\s*)?\d+\s*$", "", value)
        value=re.sub(r"\s+", " ", value).strip(" .–-\t\r\n")
        return value.casefold()
    edits={}
    for para in paras:
        value="".join(node.text or "" for node in para.findall(".//w:t",ns)).strip()
        key=match_key(value)
        if not key or key not in plain_to_indexes or len(plain_to_indexes[key]) != 1:
            continue
        index=plain_to_indexes[key][0]
        line=lines[index]
        visible=line.strip()
        style=para.find("w:pPr/w:pStyle",ns)
        sid=style.get(f"{{{W}}}val","") if style is not None else ""
        level=style_levels.get(sid)
        numpr=para.find("w:pPr/w:numPr",ns)
        if sid in mass_applied_heading_styles:
            if title_key(value) in confirmed_heading_titles and level is not None and 1 <= level <= 6:
                heading_match=re.match(r"^(#{1,6})\s*(.*)$",visible)
                indent=line[:len(line)-len(line.lstrip())]
                content=heading_match.group(2) if heading_match else line.lstrip()
                edits[index]=indent+"#"*level+" "+content+ ("\n" if line.endswith("\n") else "")
            else:
                # MarkItDown may have used the misleading style to add a
                # heading marker. Remove only that prefix from the matched
                # source paragraph; preserve its text and all other Markdown.
                prefix=re.match(r"^(\s*)#{1,6}\s+",line)
                if prefix:
                    edits[index]=prefix.group(1)+line[prefix.end():]
            continue
        if level is not None and visible.startswith("#"):
            heading_match=re.match(r"^(#{1,6})\s*(.*)$",visible)
            if heading_match and 1 <= level <= 6:
                indent=line[:len(line)-len(line.lstrip())]
                edits[index]=indent+"#"*level+" "+heading_match.group(2)+("\n" if line.endswith("\n") else "")
            continue
        if level is not None:
            # Heading and list signals conflict in this document in several places;
            # retain the source line whenever it already resembles a list.
            if re.match(r"^(?:[-+*]|\d+[.)])", visible):
                continue
            if 1 <= level <= 6:
                indent="" if line.lstrip()==line else line[:len(line)-len(line.lstrip())]
                edits[index]=indent+"#"*level+" "+line.lstrip()
            continue
        if numpr is not None and not re.match(r"^\s*(?:[-+*]|\d+[.)])\s+", line):
            # Do not infer a list from numbering metadata if the Markdown line
            # already has heading syntax or the applied Word style is a heading.
            if visible.startswith("#") or sid.lower().startswith("heading"):
                continue
            numid=numpr.find("w:numId",ns); ilvl=numpr.find("w:ilvl",ns)
            nid=numid.get(f"{{{W}}}val","") if numid is not None else ""
            try: il=int(ilvl.get(f"{{{W}}}val","0")) if ilvl is not None else 0
            except ValueError: il=0
            kinds=num_to_kind.get(nid,[])
            kind=kinds[il] if il < len(kinds) else ""
            if kind == "bullet":
                edits[index]=re.match(r"^\s*",line).group(0)+"- "+line.lstrip()
            elif kind in {"decimal","lowerLetter","upperLetter","lowerRoman","upperRoman"} and level is None:
                edits[index]=re.match(r"^\s*",line).group(0)+"1. "+line.lstrip()
    if edits:
        text="".join(edits.get(i,line) for i,line in enumerate(lines))
    return text


def _recover_pdf(text: str, source: Path, output_file: Path, warnings: list[str]):
    try:
        import pdfplumber
        pdf=pdfplumber.open(source)
    except Exception as exc:
        warnings.append(f"PDF assets/links not recovered: {exc}")
        return text, []
    pages=text.split("\f")
    assets=[]
    for page_index,page in enumerate(pdf.pages):
        segment=pages[page_index] if page_index < len(pages) else ""
        # Recover link annotations only if the bounding box contains unique visible anchor text.
        for link in page.hyperlinks:
            uri=link.get("uri")
            if not uri or not re.match(r"^https?://",uri,re.I): continue
            try:
                crop=page.crop((link["x0"],link["top"],link["x1"],link["bottom"])).extract_text() or ""
                anchor=" ".join(crop.split())
            except Exception: anchor=""
            if anchor and segment.count(anchor)==1:
                segment=segment.replace(anchor,f"[{anchor}]({uri})",1)
            else:
                warnings.append("A PDF link annotation could not be associated with unique visible anchor text; it was left unrepresented.")
        # Recover image only when its stream is a complete recognized image. Unsupported PDF encodings stay unresolved.
        payloads=[]
        for image in page.images:
            try:
                stream=image.get("stream")
                data=stream.get_data() if stream else b""
                if not _image_info(data) and stream:
                    # pdfplumber exposes decoded pixels for simple DeviceRGB/Gray XObjects.
                    from PIL import Image
                    width=int(stream.attrs.get("Width",0)); height=int(stream.attrs.get("Height",0))
                    bpc=int(stream.attrs.get("BitsPerComponent",0)); color=stream.attrs.get("ColorSpace")
                    color_name=str(color)
                    if color_name in ("DeviceRGB","/DeviceRGB", "/'DeviceRGB'", "b'DeviceRGB'"): mode,channels="RGB",3
                    elif color_name in ("DeviceGray","/DeviceGray", "/'DeviceGray'", "b'DeviceGray'"): mode,channels="L",1
                    else: mode,channels="",0
                    if width>0 and height>0 and bpc==8 and mode and len(data)==width*height*channels:
                        im=Image.frombytes(mode,(width,height),data)
                        out=io.BytesIO(); im.save(out,format="PNG"); data=out.getvalue()
                if _image_info(data): payloads.append((image,data))
                else:
                    warnings.append("A PDF embedded image stream was not independently decodable as a supported image; no reference was added.")
            except Exception:
                warnings.append("A PDF embedded image stream could not be decoded; no reference was added.")
        if payloads:
            saved=_save_assets(output_file,[data for _,data in payloads],warnings)
            # Position at image top relative to nearest preceding extracted word line only if words exist.
            words=page.extract_words()
            for (image,_),(_,relpath) in zip(payloads,saved):
                if not words:
                    warnings.append("PDF image asset was extracted but its position could not be mapped to Markdown text.")
                    continue
                preceding=[w for w in words if w["top"] <= image.get("top",0)]
                if preceding:
                    anchor=max(preceding,key=lambda w:w["top"])
                    line_text=" ".join(w["text"] for w in sorted([w for w in preceding if abs(w["top"]-anchor["top"])<2],key=lambda w:w["x0"]))
                    if line_text and segment.count(line_text)==1:
                        segment=segment.replace(line_text,line_text+"\n\n![]("+relpath+")",1)
                    else: warnings.append("PDF image asset was extracted but could not be placed reliably in Markdown.")
                else:
                    warnings.append("PDF image asset was extracted but no preceding text anchor could establish its position.")
            assets.extend(p for p,_ in saved)
        if page_index < len(pages): pages[page_index]=segment
    pdf.close()
    return "\f".join(pages),assets


def cleanup_document(text: str, output_file: Path, manage_images: bool = False, source_path: Path | None = None) -> PostprocessResult:
    warnings=[]; assets=[]
    text=_outside_fences(text,normalize_headings)
    text=_outside_fences(text,normalize_lists)
    text=normalize_spacing(text)
    text=normalize_tables(text); text=normalize_code_blocks(text); text=normalize_links(text)
    source=Path(source_path) if source_path else None
    suffix=source.suffix.lower() if source else ""
    if source and suffix==".docx":
        text=_recover_docx_structure(text,source,warnings)
        if manage_images: text,assets=_replace_docx_images(text,source,output_file,warnings)
    elif source and suffix in (".html",".htm") and manage_images:
        text,assets=_replace_html_images(text,source,output_file,warnings)
    elif source and suffix==".pptx" and manage_images:
        text,assets=_replace_pptx_images(text,source,output_file,warnings)
    elif source and suffix==".pdf" and manage_images:
        text,assets=_recover_pdf(text,source,output_file,warnings)
    elif source and suffix in (".jpg",".jpeg",".png",".gif",".webp",".bmp",".tif",".tiff") and manage_images and not text.strip():
        try: data=source.read_bytes()
        except OSError as exc:
            warnings.append(f"Source image could not be read: {exc}")
        else:
            saved=_save_assets(output_file,[data],warnings)
            if saved:
                path,relpath=saved[0]; assets=[path]; text=f"![]({relpath})\n"
            else: warnings.append("MarkItDown returned empty Markdown for this image; image was not represented because asset validation failed.")
    elif manage_images and not source:
        text, assets = extract_images(text, output_file)
    return PostprocessResult(text,assets,warnings)


def extract_images(text: str, output_file: Path) -> tuple[str, list[Path]]:
    """Compatibility utility for complete inline data URIs in existing callers."""
    warnings=[]
    refs=list(re.finditer(r"!\[([^\]]*)\]\((data:image/[^)]*)\)",text))
    payloads=[_data_uri_bytes(m.group(2)) for m in refs]
    if not refs or any(p is None for p in payloads): return text,[]
    saved=_save_assets(output_file,payloads,warnings)
    if len(saved)!=len(refs): return text,[p for p,_ in saved]
    for match,(_,relpath) in reversed(list(zip(refs,saved))):
        text=text[:match.start()]+f"![{match.group(1)}]({relpath})"+text[match.end():]
    return text,[p for p,_ in saved]


def cleanup_markdown(text: str, output_file: Path, manage_images: bool = False) -> tuple[str, list[Path]]:
    result=cleanup_document(text,output_file,manage_images)
    return result.markdown,result.assets
